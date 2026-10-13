from io import BytesIO
import math
import os

import pytest


docling = pytest.importorskip("docling")

from docx import Document

from app.ingestion.chunking import chunk_document
from app.ingestion.embeddings import ModelGatewayEmbeddingProvider, TEIEmbeddingProvider
from app.ingestion.parsers import ParserRegistry


def test_docling_parses_and_structure_chunks_docx():
    source = Document()
    source.add_heading("OAuth rollout", level=1)
    source.add_paragraph(
        "The authentication service rotates refresh tokens and verifies access tokens."
    )
    source.add_heading("Operational checklist", level=2)
    source.add_paragraph("Monitor failed callbacks and preserve workspace authorization.")
    stream = BytesIO()
    source.save(stream)

    parsed = ParserRegistry().parse(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        stream.getvalue(),
        "oauth-rollout.docx",
    )
    chunks = chunk_document(parsed, chunk_tokens=128, overlap_tokens=16)

    assert parsed.parser_name == "docling"
    assert "OAuth rollout" in parsed.pages[0].markdown_content
    assert chunks
    assert all(chunk.metadata["strategy"] == "docling_hybrid" for chunk in chunks)
    assert any("OAuth rollout" in chunk.content for chunk in chunks)


@pytest.mark.skipif(not os.getenv("TEI_RUNTIME_URL"), reason="requires a live TEI service")
def test_docling_chunks_embed_with_canonical_qwen_profile():
    source = Document()
    source.add_heading("Project knowledge", level=1)
    source.add_paragraph("Issue attachments become permission-scoped retrieval evidence.")
    stream = BytesIO()
    source.save(stream)

    parsed = ParserRegistry().parse(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        stream.getvalue(),
        "project-knowledge.docx",
    )
    chunks = chunk_document(parsed, chunk_tokens=128, overlap_tokens=16)
    provider = TEIEmbeddingProvider(
        os.environ["TEI_RUNTIME_URL"],
        "Qwen/Qwen3-Embedding-0.6B",
        "97b0c614be4d77ee51c0cef4e5f07c00f9eb65b3",
        1024,
    )

    vectors = provider.embed([chunk.content for chunk in chunks])

    assert len(vectors) == len(chunks)
    assert all(len(vector) == 1024 for vector in vectors)
    assert all(math.isclose(sum(value * value for value in vector), 1.0) for vector in vectors)


@pytest.mark.skipif(
    not os.getenv("TEI_RUNTIME_URL") or not os.getenv("MODEL_GATEWAY_RUNTIME_URL"),
    reason="requires live TEI and model gateway services",
)
def test_gateway_embedding_is_conformant_with_direct_tei():
    texts = [
        "Issue attachments become permission-scoped retrieval evidence.",
        "CommonPlan preserves citation provenance across document ingestion.",
    ]
    direct = TEIEmbeddingProvider(
        os.environ["TEI_RUNTIME_URL"],
        "Qwen/Qwen3-Embedding-0.6B",
        "97b0c614be4d77ee51c0cef4e5f07c00f9eb65b3",
        1024,
    ).embed(texts)
    gateway = ModelGatewayEmbeddingProvider(
        os.environ["MODEL_GATEWAY_RUNTIME_URL"],
        os.environ.get("MODEL_GATEWAY_API_KEY", "sk-local-commonplan-change-me"),
        "commonplan-embedding-v1",
        "Qwen/Qwen3-Embedding-0.6B",
        "97b0c614be4d77ee51c0cef4e5f07c00f9eb65b3",
        1024,
    ).embed(texts)

    assert len(gateway) == len(direct)
    for direct_vector, gateway_vector in zip(direct, gateway, strict=True):
        cosine = sum(a * b for a, b in zip(direct_vector, gateway_vector, strict=True))
        assert cosine > 0.99999
